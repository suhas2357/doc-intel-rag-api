import os
from typing import List

from langchain_core.documents import Document
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from docx import Document as DocxDocument

from app.utils.file_utils import get_file_extension


class DocumentLoaderService:
    """
    Load supported document formats into LangChain Document objects.
    Supported:
    - PDF
    - TXT
    - DOCX
    """

    def load_document(self, file_path: str) -> List[Document]:
        ext = get_file_extension(file_path)

        if ext == ".pdf":
            return self._load_pdf(file_path)
        elif ext == ".txt":
            return self._load_txt(file_path)
        elif ext == ".docx":
            return self._load_docx(file_path)
        else:
            raise ValueError(f"Unsupported document type: {ext}")

    def _load_pdf(self, file_path: str) -> List[Document]:
        """
        Load PDF using LangChain PyPDFLoader.
        Usually returns one Document per page.
        """
        loader = PyPDFLoader(file_path)
        docs = loader.load()

        for doc in docs:
            doc.metadata["file_name"] = os.path.basename(file_path)
            doc.metadata["file_type"] = ".pdf"
            doc.metadata["source"] = file_path

        return docs

    def _load_txt(self, file_path: str) -> List[Document]:
        """
        Load TXT using LangChain TextLoader.
        Usually returns one document for the full text file.
        """
        loader = TextLoader(file_path, encoding="utf-8")
        docs = loader.load()

        for doc in docs:
            doc.metadata["file_name"] = os.path.basename(file_path)
            doc.metadata["file_type"] = ".txt"
            doc.metadata["source"] = file_path

        return docs

    def _load_docx(self, file_path: str) -> List[Document]:
        """
        Load DOCX manually using python-docx.
        Convert all paragraphs into one text block document.
        """
        docx_file = DocxDocument(file_path)
        paragraphs = [p.text.strip() for p in docx_file.paragraphs if p.text.strip()]
        text = "\n".join(paragraphs).strip()

        if not text:
            return []

        return [
            Document(
                page_content=text,
                metadata={
                    "file_name": os.path.basename(file_path),
                    "file_type": ".docx",
                    "source": file_path
                }
            )
        ]